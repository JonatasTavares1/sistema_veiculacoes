import customtkinter as ctk
from tkinter import filedialog, messagebox
from controllers.veiculacao_controller import listar_veiculacoes
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import os

class ExportarView(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master)

        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("green")

        # Título
        ctk.CTkLabel(self, text="📤 Exportar Veiculações para Excel",
                     font=ctk.CTkFont(size=22, weight="bold")).pack(pady=(0, 15))

        # Botão de exportação
        ctk.CTkButton(self, text="💾 Exportar como XLSX", command=self.exportar_excel, height=40).pack(pady=10)

        # Resultado da exportação
        self.resultado_label = ctk.CTkLabel(self, text="", wraplength=500, justify="center", text_color="lightgreen")
        self.resultado_label.pack(pady=(10, 20))

        # Caixa de texto com scroll
        self.text_frame = ctk.CTkFrame(self)
        self.text_frame.pack(fill="both", expand=True)

        self.lista = ctk.CTkTextbox(self.text_frame, width=700, height=300, wrap="none", corner_radius=8)
        self.lista.pack(side="left", fill="both", expand=True)

        self.scrollbar = ctk.CTkScrollbar(self.text_frame, orientation="vertical", command=self.lista.yview)
        self.scrollbar.pack(side="right", fill="y")

        self.lista.configure(yscrollcommand=self.scrollbar.set)

        self.atualizar_lista()

    def atualizar_lista(self):
        self.lista.delete("1.0", "end")
        veiculacoes = listar_veiculacoes()

        if not veiculacoes:
            self.lista.insert("end", "Nenhuma veiculação encontrada.\n")
            return

        for v in veiculacoes:
            valor_total = (v.produto.valor_unitario * v.quantidade) - v.desconto_aplicado
            linha = (
                f"ID: {v.id} | Produto: {v.produto.nome} | PI: {v.pi.numero_pi} | "
                f"Qtd: {v.quantidade} | Desconto: R$ {v.desconto_aplicado:.2f} | "
                f"Total: R$ {valor_total:.2f} | Data: {v.data_veiculacao}\n"
            )
            self.lista.insert("end", linha)

    def exportar_excel(self):
        caminho = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if not caminho:
            return

        veiculacoes = listar_veiculacoes()
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Veiculações"

            # Cabeçalho
            headers = ["ID", "Produto", "PI", "Quantidade", "Desconto (R$)", "Valor Total (R$)", "Data"]
            ws.append(headers)

            # Estilização do cabeçalho
            for col in ws.iter_cols(min_row=1, max_row=1, min_col=1, max_col=len(headers)):
                for cell in col:
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center")

            # Dados
            for v in veiculacoes:
                valor_total = (v.produto.valor_unitario * v.quantidade) - v.desconto_aplicado
                ws.append([
                    v.id,
                    v.produto.nome,
                    v.pi.numero_pi,
                    v.quantidade,
                    round(v.desconto_aplicado, 2),
                    round(valor_total, 2),
                    v.data_veiculacao.strftime("%d/%m/%Y")
                ])

            # Autoajuste de largura
            for column_cells in ws.columns:
                max_length = max(len(str(cell.value)) for cell in column_cells)
                ws.column_dimensions[column_cells[0].column_letter].width = max_length + 2

            wb.save(caminho)
            self.resultado_label.configure(
                text=f"✅ Arquivo XLSX exportado com sucesso!\n{os.path.abspath(caminho)}"
            )
        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao exportar: {e}")
